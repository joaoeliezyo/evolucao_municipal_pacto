from io import BytesIO
from pathlib import Path
from time import time
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from sqlalchemy import delete, select

from ..models import ImportacaoPendente, Municipio, Periodo, Posicao

MESES = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO',
         'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
MUNICIPIOS = ['BETANIA', 'BURITI DOS MONTES', 'CAMPO LARGO', 'COCAL DOS ALVES',
             'CORONEL JOSÉ DIAS', 'CRISTINO CASTRO', 'IPIRANGA', 'JERUMENHA',
             'MONSENHOR GIL', 'SÃO JOSÉ DO PEIXE', 'SÃO JOSÉ DO PIAUÍ', 'TANQUE DO PIAUÍ']

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 256
MAX_ROWS = 1000
MAX_COLUMNS = 3


def validate_archive(content):
    """Reject oversized ZIP containers before openpyxl decompresses their XML."""
    if len(content) > MAX_BYTES:
        raise ValueError('O arquivo deve ter no máximo 10 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > MAX_ARCHIVE_ENTRIES or
                    sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES):
                raise ValueError('A planilha excede o limite de conteúdo descompactado.')
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError('Planilhas criptografadas não são aceitas.')
    except BadZipFile as exc:
        raise ValueError('Não foi possível abrir o arquivo Excel.') from exc


def normalize(value):
    return ' '.join(str(value or '').split()).upper()


def parse_excel(content, filename, ano, mes):
    if Path(filename).suffix.lower() != '.xlsx':
        raise ValueError('Selecione um arquivo .xlsx.')
    if not 1900 <= ano <= 2100:
        raise ValueError('Informe um ano entre 1900 e 2100.')
    if mes not in range(13):
        raise ValueError('Mês inválido.')
    validate_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('Não foi possível abrir o arquivo Excel.') from exc
    try:
        if 'Planilha1' not in workbook.sheetnames:
            raise ValueError('A aba Planilha1 não foi encontrada.')
        sheet = workbook['Planilha1']
        if (sheet.max_row and sheet.max_row > MAX_ROWS) or (sheet.max_column and sheet.max_column > MAX_COLUMNS):
            raise ValueError('A planilha deve ter até 1000 linhas e exatamente três colunas.')
        # Ignore untrusted worksheet dimensions and bound actual XML iteration too.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS + 1, values_only=True)
        header = next(rows, ())
        if [normalize(v) for v in header[:3]] != ['MUNICIPIO', 'POSIÇÃO', 'MÊS'] or header[3] is not None:
            raise ValueError('Cabeçalhos esperados: MUNICIPIO, POSIÇÃO e MÊS, nessa ordem.')
        result, seen = [], set()
        for number, row in enumerate(rows, 2):
            if all(v is None for v in row):
                continue
            if number > MAX_ROWS or row[3] is not None:
                raise ValueError('A planilha deve ter até 1000 linhas e exatamente três colunas.')
            original, position, month = row[:3]
            name = normalize(original)
            if name not in MUNICIPIOS:
                raise ValueError(f'Linha {number}: município desconhecido ou vazio.')
            if isinstance(position, bool) or not isinstance(position, (int, float)) or position <= 0 or position != int(position):
                raise ValueError(f'Linha {number}: posição deve ser um inteiro positivo.')
            if position > 2**63 - 1:
                raise ValueError(f'Linha {number}: posição excede o limite permitido.')
            month = normalize(month)
            if month not in MESES:
                raise ValueError(f'Linha {number}: mês inválido.')
            month_number = MESES.index(month) + 1
            if mes and mes != month_number:
                raise ValueError('O mês selecionado não coincide com todos os registros. Para arquivos com vários meses, selecione Todos os meses.')
            key = (name, month_number)
            if key in seen:
                raise ValueError(f'Linha {number}: município duplicado no mesmo mês.')
            seen.add(key)
            result.append({'original': str(original), 'nome': name, 'posicao': int(position), 'mes': month_number})
        if not result:
            raise ValueError('A planilha não contém registros.')
        return result
    except (DefusedXmlException, ParseError, BadZipFile) as exc:
        raise ValueError('Não foi possível abrir o arquivo Excel com segurança.') from exc
    finally:
        workbook.close()


def conflicts(session, ano, records):
    months = {r['mes'] for r in records}
    return list(session.scalars(select(Periodo.mes).where(Periodo.ano == ano, Periodo.mes.in_(months))))


def persist(session, ano, filename, records, replace=False, pending_token=None):
    # Validation and all writes share one transaction, including replacement.
    with session.begin():
        if pending_token is not None:
            consumed = session.execute(delete(ImportacaoPendente).where(
                ImportacaoPendente.token == pending_token,
                ImportacaoPendente.expira_em > int(time()),
            ))
            if consumed.rowcount != 1:
                raise ValueError('A revisão expirou ou já foi confirmada. Selecione o arquivo e valide novamente.')
        if conflicts(session, ano, records) and not replace:
            raise ValueError('Há competências existentes. Confirme a substituição para continuar.')
        for month in sorted({r['mes'] for r in records}):
            period = session.scalar(select(Periodo).where(Periodo.ano == ano, Periodo.mes == month))
            if period:
                session.execute(delete(Posicao).where(Posicao.periodo_id == period.id))
                session.delete(period)
                session.flush()
            period = Periodo(ano=ano, mes=month, arquivo_nome=filename)
            session.add(period)
            session.flush()
            for record in (r for r in records if r['mes'] == month):
                city = session.scalar(select(Municipio).where(Municipio.nome_exibicao == record['nome']))
                if city is None:
                    city = Municipio(nome_original=record['original'], nome_exibicao=record['nome'])
                    session.add(city)
                    session.flush()
                session.add(Posicao(municipio_id=city.id, periodo_id=period.id, posicao=record['posicao']))
